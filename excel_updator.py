import sys
import logging
import traceback
from pathlib import Path
from openpyxl import load_workbook

# ─────── Logging setup ───────
LOG_DIR = Path.home() / "ReceiptsLogs"
LOG_DIR.mkdir(exist_ok=True)

LOG_FILE = LOG_DIR / "receipt_copy.log"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)

log = logging.getLogger("SheetCopy")
# ────────────────────────────


def copy_sheets_by_index(source_file, target_file):
    log.info(f"Starting copy")
    log.info(f"Source file: {source_file}")
    log.info(f"Target file: {target_file}")

    try:
        source_path = Path(source_file)
        target_path = Path(target_file)

        if not source_path.exists():
            raise FileNotFoundError(f"Source file not found: {source_path}")

        if not target_path.exists():
            raise FileNotFoundError(f"Target file not found: {target_path}")

        log.info("Loading source workbook")
        src_wb = load_workbook(source_path)

        log.info("Loading target workbook")
        tgt_wb = load_workbook(target_path)

        src_sheets = src_wb.worksheets
        tgt_sheets = tgt_wb.worksheets

        count = min(len(src_sheets), len(tgt_sheets))
        log.info(f"Matching {count} sheet(s) by index")

        print(f"\n📁 Copying data from '{source_file}' to '{target_file}'")
        print(f"➡️  Matching {count} sheet(s) by index...")

        for i in range(count):
            src_ws = src_sheets[i]
            tgt_ws = tgt_sheets[i]

            log.info(f"Copying sheet {i+1}: {src_ws.title} → {tgt_ws.title}")
            print(f"  Sheet {i+1}: '{src_ws.title}' → '{tgt_ws.title}'")

            for row in src_ws.iter_rows():
                for cell in row:
                    tgt_cell = tgt_ws[cell.coordinate]
                    tgt_cell.value = cell.value
                    tgt_cell.number_format = cell.number_format

        log.info("Saving target workbook")
        tgt_wb.save(target_path)

        print(f"✅ Saved: {target_file}")
        log.info("Copy completed successfully")

    except Exception as e:
        log.error("Copy failed")
        log.error(str(e))
        log.error(traceback.format_exc())
        print("❌ Error occurred. See log for details:")
        print(LOG_FILE)


def copy_multiple_pairs(file_pairs):
    log.info("Starting batch copy")
    for source_file, target_file in file_pairs:
        copy_sheets_by_index(source_file, target_file)


if __name__ == "__main__":
    log.info("Script started")

    try:
        file_pairs = [
            (
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\February.xlsx",
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\February_2025.xlsx"
            ),
            (
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\March.xlsx",
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\March_2025.xlsx"
            ),
            (
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\April.xlsx",
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\April_2025.xlsx"
            ),
            (
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\May.xlsx",
                r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts\May_2025.xlsx"
            )
        ]

        copy_multiple_pairs(file_pairs)

    except Exception:
        log.critical("Fatal error in main execution")
        log.critical(traceback.format_exc())
