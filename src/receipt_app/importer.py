from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from dateutil import parser as date_parser
from openpyxl import load_workbook

from .config import CORE_FIELD_KEYS, FieldDefinition
from .database import ReceiptDatabase, ReceiptRecord


@dataclass
class ImportResult:
    imported: int = 0
    skipped: int = 0
    errors: list[str] | None = None

    def add_error(self, message: str) -> None:
        if self.errors is None:
            self.errors = []
        self.errors.append(message)


FIELD_ALIASES: dict[str, set[str]] = {
    "receipt_no": {"receipt no", "receipt number", "receipt", "receipt_no", "receiptno"},
    "transaction_date": {"date", "transaction date", "payment date", "transaction_date"},
    "name": {"name", "customer", "customer name", "description", "member name"},
    "amount": {"amount", "payment amount", "total", "value", "payment_amount"},
    "payment_type": {"payment type", "method", "payment method", "type", "payment_type"},
    "member_no": {"member no", "member number", "member", "member_no", "memberno"},
    "notes": {"notes", "note", "comments", "comment"},
}


def normalise_header(header: Any) -> str:
    return re.sub(r"[\W_]+", " ", str(header or "").strip().lower()).strip()


def normalise_amount(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    if cleaned in {"", ".", "-"}:
        return None
    return float(cleaned)


def normalise_date(value: Any, fallback_sheet_name: str | None = None) -> str:
    raw = value if value not in {None, ""} else fallback_sheet_name
    if raw in {None, ""}:
        return ""
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    try:
        return date_parser.parse(str(raw), dayfirst=True, fuzzy=True).date().isoformat()
    except Exception:
        return str(raw).strip()


def infer_mapping(headers: list[Any], fields: list[FieldDefinition]) -> dict[str, int]:
    cleaned_headers = [normalise_header(header) for header in headers]
    mapping: dict[str, int] = {}
    for field_def in fields:
        aliases = FIELD_ALIASES.get(field_def.key, {field_def.label.lower(), field_def.key.lower()})
        aliases = {normalise_header(alias) for alias in aliases | {field_def.label, field_def.key}}
        for index, header in enumerate(cleaned_headers):
            if header in aliases:
                mapping[field_def.key] = index
                break
    return mapping


def row_to_record(row: list[Any], mapping: dict[str, int], fields: list[FieldDefinition], source: str, sheet_name: str | None = None) -> ReceiptRecord:
    data: dict[str, Any] = {}
    for field_def in fields:
        index = mapping.get(field_def.key)
        data[field_def.key] = row[index] if index is not None and index < len(row) else ""

    custom = {
        field_def.key: str(data.get(field_def.key, "") or "").strip()
        for field_def in fields
        if field_def.key not in CORE_FIELD_KEYS and str(data.get(field_def.key, "") or "").strip()
    }
    return ReceiptRecord(
        receipt_no=str(data.get("receipt_no", "") or "").strip(),
        transaction_date=normalise_date(data.get("transaction_date"), sheet_name),
        name=str(data.get("name", "") or "").strip(),
        amount=normalise_amount(data.get("amount")),
        payment_type=str(data.get("payment_type", "") or "").strip(),
        member_no=str(data.get("member_no", "") or "").strip(),
        notes=str(data.get("notes", "") or "").strip(),
        custom_fields=custom,
        source=source,
    )


def has_meaningful_data(record: ReceiptRecord) -> bool:
    return any([record.receipt_no, record.name, record.amount is not None, record.member_no, record.notes, record.custom_fields])


def import_csv(path: Path, db: ReceiptDatabase, fields: list[FieldDefinition]) -> ImportResult:
    result = ImportResult(errors=[])
    with Path(path).open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        headers = next(reader, None)
        if not headers:
            result.add_error("The CSV file appears to be empty.")
            return result
        mapping = infer_mapping(headers, fields)
        for row_number, row in enumerate(reader, start=2):
            try:
                record = row_to_record(row, mapping, fields, f"csv:{Path(path).name}")
                if not has_meaningful_data(record):
                    result.skipped += 1
                    continue
                db.upsert_receipt(record)
                result.imported += 1
            except Exception as exc:
                result.skipped += 1
                result.add_error(f"Row {row_number}: {exc}")
    return result


def import_excel(path: Path, db: ReceiptDatabase, fields: list[FieldDefinition]) -> ImportResult:
    result = ImportResult(errors=[])
    workbook = load_workbook(path, data_only=True)
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = list(rows[0])
        mapping = infer_mapping(headers, fields)
        if not mapping:
            result.add_error(f"Sheet '{sheet_name}' was skipped because no matching receipt columns were found.")
            continue
        for row_number, raw_row in enumerate(rows[1:], start=2):
            try:
                record = row_to_record(list(raw_row), mapping, fields, f"excel:{Path(path).name}:{sheet_name}", sheet_name)
                if not has_meaningful_data(record):
                    result.skipped += 1
                    continue
                db.upsert_receipt(record)
                result.imported += 1
            except Exception as exc:
                result.skipped += 1
                result.add_error(f"Sheet {sheet_name}, row {row_number}: {exc}")
    return result


def import_file(path: Path, db: ReceiptDatabase, fields: list[FieldDefinition]) -> ImportResult:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return import_csv(path, db, fields)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return import_excel(path, db, fields)
    result = ImportResult(errors=[])
    result.add_error("Unsupported file type. Please upload a CSV or Excel workbook.")
    return result


def mapping_preview(path: Path, fields: list[FieldDefinition]) -> dict[str, Any]:
    suffix = Path(path).suffix.lower()
    headers: list[Any] = []
    sample_rows: list[list[Any]] = []
    if suffix == ".csv":
        with Path(path).open("r", newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            headers = next(reader, [])
            sample_rows = [row for _, row in zip(range(5), reader)]
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        rows = list(worksheet.iter_rows(values_only=True, max_row=6))
        if rows:
            headers = list(rows[0])
            sample_rows = [list(row) for row in rows[1:]]
    mapping = infer_mapping(headers, fields)
    return {
        "headers": [str(h or "") for h in headers],
        "mapping": mapping,
        "sample_rows": sample_rows,
        "sample_json": json.dumps(sample_rows[:3], default=str, indent=2),
    }
