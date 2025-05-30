import calendar
from datetime import datetime
from openpyxl import Workbook
import os

def create_monthly_workbooks(year, output_dir):
    os.makedirs(output_dir, exist_ok=True)  # Ensure the output directory exists

    for month in range(1, 13):  # January to December
        month_name = calendar.month_name[month]
        wb = Workbook()
        
        # Remove default sheet
        default_sheet = wb.active
        wb.remove(default_sheet)
        
        num_days = calendar.monthrange(year, month)[1]
        
        for day in range(1, num_days + 1):
            # Use "-" instead of "/" to avoid invalid sheet name
            sheet_name = f"{day:02d}-{month:02d}-{year}"
            ws = wb.create_sheet(title=sheet_name)
            ws["A1"] = f"Sheet for {sheet_name}"

        file_name = f"{month_name}_{year}.xlsx"
        file_path = os.path.join(output_dir, file_name)
        wb.save(file_path)
        print(f"Created: {file_path}")

if __name__ == "__main__":
    year = 2025
    output_directory = r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Apps\x"  # Adjust path if needed
    create_monthly_workbooks(year, output_directory)
