import os
import re
from datetime import datetime
import pandas as pd
from docx import Document
from openpyxl import load_workbook

# CONFIG - change these paths
INPUT_FOLDER = r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\OldWordFiles"
OUTPUT_FOLDER = r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts"

EXPECTED_COLUMNS = ['Description', 'Payment amount', 'Payment Type', 'Member No.', 'Receipt No.', 'Notes']

def day_suffix(day: int) -> str:
    if 11 <= day <= 13:
        return "th"
    return {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")

def extract_date(filename: str):
    # Matches DD.MM.YYYY, e.g. 21.05.2025
    m = re.search(r'(\d{2}\.\d{2}\.\d{4})', filename)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%d.%m.%Y").date()
    except ValueError:
        return None

def read_docx_table_to_df(docx_path):
    doc = Document(docx_path)
    if not doc.tables:
        print(f"⚠️ No tables found in '{docx_path}'")
        return None
    
    table = doc.tables[0]

    headers = [cell.text.strip() for cell in table.rows[0].cells]

    # Map expected columns to table columns by case-insensitive match
    col_indices = {}
    for expected_col in EXPECTED_COLUMNS:
        for i, header in enumerate(headers):
            if header.lower() == expected_col.lower():
                col_indices[expected_col] = i
                break

    missing = set(EXPECTED_COLUMNS) - set(col_indices.keys())
    if missing:
        print(f"⚠️ Missing columns {missing} in '{docx_path}'")
        return None

    data = []
    for row in table.rows[1:]:
        row_data = []
        for col in EXPECTED_COLUMNS:
            row_data.append(row.cells[col_indices[col]].text.strip())
        data.append(row_data)

    df = pd.DataFrame(data, columns=EXPECTED_COLUMNS)
    # Clean Payment amount column
    df['Payment amount'] = df['Payment amount'].str.replace('[$,]', '', regex=True).astype(float)
    return df

def save_to_monthly_workbook(df, date, output_folder):
    month_str = date.strftime("%B")
    day_num = date.day
    suffix = day_suffix(day_num)
    day_sheet = f"{day_num}{suffix}"

    os.makedirs(output_folder, exist_ok=True)
    excel_path = os.path.join(output_folder, f"{month_str}.xlsx")

    if os.path.exists(excel_path):
        book = load_workbook(excel_path)
        if day_sheet in book.sheetnames:
            # Remove existing sheet so we can replace it
            std = book[day_sheet]
            book.remove(std)
        with pd.ExcelWriter(excel_path, engine='openpyxl', mode='a') as writer:
            writer.book = book
            df.to_excel(writer, sheet_name=day_sheet, index=False)
            writer.save()
    else:
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name=day_sheet, index=False)

    print(f"✅ Saved {len(df)} rows to {excel_path} sheet '{day_sheet}'")

def main():
    files = [f for f in os.listdir(INPUT_FOLDER) if f.lower().endswith('.docx')]
    if not files:
        print(f"❌ No .docx files found in {INPUT_FOLDER}")
        return

    for file in files:
        path = os.path.join(INPUT_FOLDER, file)
        date = extract_date(file)
        if not date:
            print(f"⚠️ Skipping '{file}': could not extract valid date.")
            continue

        df = read_docx_table_to_df(path)
        if df is None:
            print(f"⚠️ Skipping '{file}': failed to read table.")
            continue

        save_to_monthly_workbook(df, date, OUTPUT_FOLDER)

if __name__ == "__main__":
    main()
