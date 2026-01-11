import os
import re
import sys
import csv
import logging
import traceback
import sqlite3
import tkinter as tk
from pathlib import Path
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook
from dateutil.parser import parse

# ─────── Logging setup ───────
LOG_DIR = Path.home() / "ReceiptsLogs"
LOG_DIR.mkdir(exist_ok=True)

LOG_FILE = LOG_DIR / "receipt_query_app.log"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)

log = logging.getLogger("ReceiptQueryApp")
# ────────────────────────────

# ─────── Configuration ───────
DEFAULT_INPUT_DIR = r"C:\Users\reception\OneDrive - Coogee Legion Ex-Services Club\Documents\Receipts"
DEFAULT_OUTPUT_DB = os.path.join(DEFAULT_INPUT_DIR, "transactions.db")
# ────────────────────────────


def extract_and_insert(excel_dir: str, db_path: str):
    """
    Scan `excel_dir` recursively for Excel files, rebuild SQLite DB,
    and insert parsed transactions.
    """
    log.info("Starting database rebuild")
    log.info(f"Excel directory: {excel_dir}")
    log.info(f"Database path: {db_path}")
    log.info(f"Working directory: {os.getcwd()}")

    expected = {"description", "payment amount", "payment type", "member no", "receipt no", "notes"}

    if not os.path.exists(excel_dir):
        raise FileNotFoundError(f"Input directory does not exist: {excel_dir}")

    os.makedirs(os.path.dirname(db_path), exist_ok=True)

    try:
        with sqlite3.connect(db_path) as conn:
            log.info("Connected to SQLite")

            conn.execute("DROP TABLE IF EXISTS transactions;")
            conn.execute(
                """
                CREATE TABLE transactions (
                    description      TEXT,
                    payment_amount   REAL,
                    payment_type     TEXT,
                    member_no        TEXT,
                    receipt_no       TEXT,
                    notes            TEXT,
                    transaction_date TEXT
                );
                """
            )
            log.info("Recreated transactions table")

            for root, _, files in os.walk(excel_dir):
                for fname in files:
                    if not fname.lower().endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
                        continue

                    fp = os.path.join(root, fname)
                    log.info(f"Processing file: {fp}")

                    try:
                        wb = load_workbook(fp, data_only=True)
                    except Exception as e:
                        log.warning(f"Skipping unreadable file: {fp}")
                        log.warning(str(e))
                        continue

                    for sheet_name in wb.sheetnames:
                        ws = wb[sheet_name]
                        log.info(f"Parsing sheet: {sheet_name}")

                        try:
                            header = list(next(ws.iter_rows(values_only=True)))
                        except Exception as e:
                            log.warning(f"Failed reading header in {fp} [{sheet_name}]")
                            continue

                        cols = [
                            re.sub(r"[\W_]+", " ", str(h).strip().lower()).strip()
                            for h in header
                        ]

                        if not expected.issubset(cols):
                            log.info(f"Skipping sheet due to missing columns: {sheet_name}")
                            continue

                        db_cols = [c.replace(' ', '_') for c in cols] + ['transaction_date']
                        placeholders = ",".join("?" for _ in db_cols)
                        sql = f"INSERT INTO transactions ({','.join(db_cols)}) VALUES ({placeholders})"

                        for row in ws.iter_rows(min_row=2, values_only=True):
                            row_dict = dict(zip(cols, row))

                            try:
                                date_obj = parse(str(sheet_name), dayfirst=True)
                                date_str = date_obj.strftime('%Y-%m-%d')
                            except Exception:
                                date_str = sheet_name

                            values = [row_dict[c] for c in cols] + [date_str]

                            try:
                                conn.execute(sql, values)
                            except Exception as e:
                                log.error("Failed inserting row")
                                log.error(f"File: {fp} | Sheet: {sheet_name}")
                                log.error(str(e))
                                log.error(traceback.format_exc())

            conn.commit()
            log.info("Database rebuild complete")

    except Exception:
        log.critical("Database rebuild failed")
        log.critical(traceback.format_exc())
        raise


class App(tk.Tk):
    def __init__(self):
        super().__init__()

        log.info("Application starting")

        self.title("Receipt Query App")
        self.geometry("980x680")
        self.input_dir = DEFAULT_INPUT_DIR
        self.db_path = DEFAULT_OUTPUT_DB

        # Directory selector
        dir_frame = ttk.Frame(self)
        dir_frame.pack(fill='x', padx=10, pady=5)
        self.dir_label = ttk.Label(dir_frame, text=f"Input folder: {self.input_dir}")
        self.dir_label.pack(side='left', expand=True)
        ttk.Button(dir_frame, text="Change Folder", command=self.select_folder).pack(side='right')

        # Controls
        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill='x', padx=10, pady=5)
        for txt, cmd in [
            ("Recreate & Refresh DB", self.refresh_db),
            ("Show All", self.show_all),
            ("Export CSV", self.export_csv),
        ]:
            ttk.Button(btn_frame, text=txt, command=cmd).pack(side='left', padx=(0,5))

        self._add_search_field("Receipt No:", "receipt_no", self.search_receipt)
        self._add_search_field("Member No:", "member_no", self.search_member)
        self._add_search_field("Payment Amount:", "payment_amount", self.search_amount)

        # Date search
        date_frame = ttk.Frame(self)
        date_frame.pack(fill='x', padx=10, pady=2)
        ttk.Label(date_frame, text="Transaction Date (DD/MM/YYYY):").pack(side='left')
        self.day_var = tk.StringVar(); ttk.Entry(date_frame, textvariable=self.day_var, width=4).pack(side='left', padx=(5,2))
        ttk.Label(date_frame, text="/").pack(side='left')
        self.month_var = tk.StringVar(); ttk.Entry(date_frame, textvariable=self.month_var, width=4).pack(side='left', padx=2)
        ttk.Label(date_frame, text="/").pack(side='left')
        self.year_var = tk.StringVar(); ttk.Entry(date_frame, textvariable=self.year_var, width=6).pack(side='left', padx=(2,10))
        ttk.Button(date_frame, text="Search", command=self.search_date).pack(side='left')

        # Results table
        self.tree = ttk.Treeview(self, show='headings')
        ttk.Scrollbar(self, orient='vertical', command=self.tree.yview).pack(side='right', fill='y')
        ttk.Scrollbar(self, orient='horizontal', command=self.tree.xview).pack(side='bottom', fill='x')
        self.tree.pack(fill='both', expand=True, padx=10, pady=(0,10))

        self.last_cols = []
        self.last_rows = []

    def _add_search_field(self, label, col, fn):
        f = ttk.Frame(self); f.pack(fill='x', padx=10, pady=2)
        ttk.Label(f, text=label).pack(side='left')
        var = tk.StringVar()
        ttk.Entry(f, textvariable=var, width=30).pack(side='left', padx=5)
        ttk.Button(f, text="Search", command=lambda: fn(var.get().strip())).pack(side='left')
        setattr(self, f"{col}_var", var)

    def select_folder(self):
        folder = filedialog.askdirectory(initialdir=os.path.expanduser('~'))
        if folder:
            log.info(f"User selected new folder: {folder}")
            self.input_dir = folder
            self.db_path = os.path.join(folder, "transactions.db")
            self.dir_label['text'] = f"Input folder: {folder}"

    def refresh_db(self):
        try:
            extract_and_insert(self.input_dir, self.db_path)
            messagebox.showinfo("Success", "Database recreated.")
        except Exception:
            messagebox.showerror("Error", f"Failed to recreate DB.\nSee log:\n{LOG_FILE}")

    def show_all(self):
        self._run_search("1=1", ())

    def export_csv(self):
        if not self.last_rows:
            return messagebox.showwarning("No Data", "No results to export.")
        path = filedialog.asksaveasfilename(defaultextension='.csv', filetypes=[("CSV","*.csv")])
        if path:
            try:
                with open(path, 'w', newline='', encoding='utf-8') as f:
                    w = csv.writer(f)
                    w.writerow(self.last_cols)
                    w.writerows(self.last_rows)
                log.info(f"Exported CSV to {path}")
                messagebox.showinfo("Exported", f"Saved to {path}")
            except Exception:
                log.error("CSV export failed")
                log.error(traceback.format_exc())
                messagebox.showerror("Error", "Failed to export CSV")

    def search_receipt(self, r):
        if not r:
            return messagebox.showwarning("Input","Enter Receipt No.")
        self._run_search("receipt_no = ?", (r,))

    def search_member(self, m):
        if not m:
            return messagebox.showwarning("Input","Enter Member No.")
        self._run_search("member_no = ?", (m,))

    def search_amount(self, a):
        try:
            a = float(a)
        except:
            return messagebox.showwarning("Input","Enter numeric Amount.")
        self._run_search("payment_amount = ?", (a,))

    def search_date(self):
        d = self.day_var.get().zfill(2)
        m = self.month_var.get().zfill(2)
        y = self.year_var.get().zfill(4)
        if not (d.isdigit() and m.isdigit() and y.isdigit()):
            return messagebox.showwarning("Input","Enter numeric date.")
        self._run_search("transaction_date = ?", (f"{y}-{m}-{d}",))

    def _run_search(self, where, params):
        log.info(f"Running query WHERE {where} PARAMS {params}")
        try:
            with sqlite3.connect(self.db_path) as conn:
                cur = conn.cursor()
                cur.execute(f"SELECT * FROM transactions WHERE {where};", params)
                rows = cur.fetchall()
                cols = [d[0] for d in cur.description]
        except Exception:
            log.error("Query failed")
            log.error(traceback.format_exc())
            return messagebox.showerror("Search Error", "Query failed. See log.")

        if not rows:
            return messagebox.showinfo("No Results","No matching transactions.")

        self.last_cols, self.last_rows = cols, rows
        self.display_rows(rows, cols)

    def display_rows(self, rows, cols):
        self.tree.delete(*self.tree.get_children())
        self.tree['columns'] = cols
        for c in cols:
            self.tree.heading(c, text=c)
            self.tree.column(c, width=120, anchor='w')
        for r in rows:
            self.tree.insert('', 'end', values=r)


if __name__ == '__main__':
    try:
        App().mainloop()
    except Exception:
        log.critical("Fatal application crash")
        log.critical(traceback.format_exc())
